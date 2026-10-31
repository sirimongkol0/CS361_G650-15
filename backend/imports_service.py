"""Validation and background processing for V4-4 data imports.

The uploaded source file is temporary: it is written to the configured object
storage, parsed in a background task, and deleted when validation finishes.
Only whitelisted, non-identity fields are kept in ImportBatch.preview_data.
"""

import csv
import io
import os
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

import models
import storage
from database import SessionLocal


FEEDBACK_COLUMNS = (
    "external_id", "title", "respondent_type_code", "date", "rating", "comment",
    "partner_id", "activity_id", "document_id", "exchange_id",
)
OUTCOME_COLUMNS = (
    "external_id", "outcome_type_code", "title", "description", "value", "unit",
    "occurred_on", "evidence_url", "partner_id", "activity_id", "document_id", "exchange_id",
)

REQUIRED_COLUMNS = {
    "feedback": ("external_id", "title", "respondent_type_code"),
    "outcome": ("external_id", "outcome_type_code", "title"),
}

CONTEXT_MODELS = {
    "partner_id": models.Partner,
    "activity_id": models.Activity,
    "document_id": models.Document,
    "exchange_id": models.ExchangeStudent,
}


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _read_rows(file_bytes: bytes, filename: str) -> list[tuple[int, dict[str, str]]]:
    extension = os.path.splitext(filename.lower())[1]
    if extension == ".csv":
        try:
            text = file_bytes.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("ไฟล์ CSV ต้องบันทึกเป็น UTF-8 (แนะนำ UTF-8 with BOM)") from exc
        reader = csv.reader(io.StringIO(text))
        headers = next(reader, None)
        if not headers:
            raise ValueError("ไม่พบแถวหัวตารางในไฟล์ CSV")
        headers = [str(name or "").strip().lower() for name in headers]
        result = []
        previous_line = reader.line_num
        for values in reader:
            row_number = previous_line + 1
            previous_line = reader.line_num
            if not any(_text(value) for value in values):
                continue
            result.append((row_number, {
                headers[i]: _text(values[i] if i < len(values) else None)
                for i in range(len(headers))
            }))
        return result

    if extension == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("เซิร์ฟเวอร์ยังไม่ได้ติดตั้ง openpyxl กรุณาติดตั้งตาม requirements.txt") from exc
        try:
            workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
            sheet = workbook.active
            iterator = sheet.iter_rows(values_only=True)
            header_row = next(iterator, None)
            if not header_row:
                raise ValueError("ไม่พบแถวหัวตารางในไฟล์ XLSX")
            headers = [(_text(value).lower()) for value in header_row]
            result = []
            for row_number, values in enumerate(iterator, start=2):
                if not any(_text(v) for v in values):
                    continue
                result.append((row_number, {
                    headers[i]: _text(values[i] if i < len(values) else None)
                    for i in range(len(headers)) if headers[i]
                }))
            workbook.close()
            return result
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("ไม่สามารถอ่านไฟล์ XLSX ได้ กรุณาตรวจสอบว่าไฟล์ไม่เสียหาย") from exc

    raise ValueError("รองรับเฉพาะไฟล์ .csv และ .xlsx")


def _parse_int(value: str, field: str, errors: list[str], *, minimum: int | None = None,
               maximum: int | None = None) -> int | None:
    if not value:
        return None
    try:
        parsed = int(value)
    except (TypeError, ValueError):
        errors.append(f"{field} ต้องเป็นจำนวนเต็ม")
        return None
    if minimum is not None and parsed < minimum:
        errors.append(f"{field} ต้องไม่น้อยกว่า {minimum}")
    if maximum is not None and parsed > maximum:
        errors.append(f"{field} ต้องไม่เกิน {maximum}")
    return parsed


def _parse_date(value: str, field: str, errors: list[str]) -> str | None:
    if not value:
        return None
    try:
        return date.fromisoformat(value[:10]).isoformat()
    except ValueError:
        errors.append(f"{field} ต้องอยู่ในรูปแบบ YYYY-MM-DD")
        return None


def _parse_context(raw: dict[str, str], db, errors: list[str]) -> dict[str, int | None]:
    context: dict[str, int | None] = {}
    for field, model in CONTEXT_MODELS.items():
        value = _parse_int(raw.get(field, ""), field, errors, minimum=1)
        context[field] = value
        if value is None:
            continue
        row = db.get(model, value)
        if row is None or getattr(row, "is_deleted", False):
            errors.append(f"ไม่พบข้อมูล {field} หรือรายการถูกลบแล้ว")
    if not any(context.values()):
        errors.append("ต้องระบุบริบทอย่างน้อย 1 รายการ: partner_id, activity_id, document_id หรือ exchange_id")
    return context


def _validate_feedback(raw: dict[str, str], db, source_id: int, seen_ids: set[str]) -> tuple[dict | None, list[str]]:
    errors: list[str] = []
    external_id = raw.get("external_id", "").strip()
    title = raw.get("title", "").strip()
    respondent_code = raw.get("respondent_type_code", "").strip()
    if not external_id:
        errors.append("ต้องระบุ external_id")
    elif len(external_id) > 255:
        errors.append("external_id ต้องมีความยาวไม่เกิน 255 ตัวอักษร")
    elif external_id in seen_ids:
        errors.append("external_id ซ้ำภายในไฟล์")
    else:
        # Count every occurrence, even when another field on this row is invalid.
        seen_ids.add(external_id)
        if db.query(models.Feedback.id).filter_by(feedback_source_id=source_id, external_id=external_id).first():
            errors.append("external_id นี้มีอยู่ในระบบแล้ว")
    if not title:
        errors.append("ต้องระบุ title")
    elif len(title) > 255:
        errors.append("title ต้องมีความยาวไม่เกิน 255 ตัวอักษร")
    respondent = db.query(models.RespondentType).filter_by(code=respondent_code, is_active=True).first()
    if not respondent:
        errors.append("ไม่พบ respondent_type_code ที่เปิดใช้งาน")

    context = _parse_context(raw, db, errors)
    rating = _parse_int(raw.get("rating", ""), "rating", errors, minimum=1, maximum=5)
    feedback_date = _parse_date(raw.get("date", ""), "date", errors)
    comment = raw.get("comment", "").strip() or None
    if comment and len(comment) > 4000:
        errors.append("comment ต้องมีความยาวไม่เกิน 4000 ตัวอักษร")

    if errors:
        return None, errors
    return {
        "external_id": external_id,
        "title": title,
        "respondent_type_id": respondent.id,
        "date": feedback_date,
        "rating": rating,
        "comment": comment,
        **context,
    }, []


def _validate_outcome(raw: dict[str, str], db, seen_ids: set[str]) -> tuple[dict | None, list[str]]:
    errors: list[str] = []
    external_id = raw.get("external_id", "").strip()
    title = raw.get("title", "").strip()
    type_code = raw.get("outcome_type_code", "").strip()
    if not external_id:
        errors.append("ต้องระบุ external_id")
    elif len(external_id) > 255:
        errors.append("external_id ต้องมีความยาวไม่เกิน 255 ตัวอักษร")
    elif external_id in seen_ids:
        errors.append("external_id ซ้ำภายในไฟล์")
    else:
        # Count every occurrence, even when another field on this row is invalid.
        seen_ids.add(external_id)
        if db.query(models.Outcome.id).filter_by(external_id=external_id).first():
            errors.append("external_id นี้มีอยู่ในระบบแล้ว")
    if not title:
        errors.append("ต้องระบุ title")
    elif len(title) > 255:
        errors.append("title ต้องมีความยาวไม่เกิน 255 ตัวอักษร")
    outcome_type = db.query(models.OutcomeType).filter_by(code=type_code, is_active=True).first()
    if not outcome_type:
        errors.append("ไม่พบ outcome_type_code ที่เปิดใช้งาน")
    context = _parse_context(raw, db, errors)
    occurred_on = _parse_date(raw.get("occurred_on", ""), "occurred_on", errors)
    value_text = raw.get("value", "").strip()
    value = None
    if value_text:
        try:
            number = Decimal(value_text)
            if not number.is_finite() or number < 0:
                raise InvalidOperation
            rounded = number.quantize(Decimal("0.01"))
            if rounded != number:
                errors.append("value ต้องมีทศนิยมไม่เกิน 2 ตำแหน่ง")
            else:
                value = str(rounded)
        except (InvalidOperation, ValueError):
            errors.append("value ต้องเป็นตัวเลขตั้งแต่ 0 ขึ้นไป และมีทศนิยมไม่เกิน 2 ตำแหน่ง")
    evidence_url = raw.get("evidence_url", "").strip() or None
    if evidence_url and len(evidence_url) > 2000:
        errors.append("evidence_url ต้องมีความยาวไม่เกิน 2000 ตัวอักษร")
    description = raw.get("description", "").strip() or None
    if description and len(description) > 4000:
        errors.append("description ต้องมีความยาวไม่เกิน 4000 ตัวอักษร")
    unit = raw.get("unit", "").strip()
    if not unit and outcome_type:
        unit = outcome_type.unit
    unit = unit or None
    if unit and len(unit) > 100:
        errors.append("unit ต้องมีความยาวไม่เกิน 100 ตัวอักษร")

    if errors:
        return None, errors
    return {
        "external_id": external_id,
        "outcome_type_id": outcome_type.id,
        "title": title,
        "description": description,
        "value": value,
        "unit": unit,
        "occurred_on": occurred_on,
        "evidence_url": evidence_url,
        **context,
    }, []


def validate_import(file_bytes: bytes, filename: str, kind: str, db,
                    source_id: int) -> tuple[list[dict], list[dict], int]:
    """Return (valid normalized rows, row errors, total data rows)."""
    if kind not in REQUIRED_COLUMNS:
        raise ValueError("ประเภทข้อมูลนำเข้าไม่ถูกต้อง")
    rows = _read_rows(file_bytes, filename)
    if not rows:
        raise ValueError("ไฟล์ไม่มีข้อมูลสำหรับนำเข้า")
    present = set(rows[0][1].keys())
    missing = [column for column in REQUIRED_COLUMNS[kind] if column not in present]
    if missing:
        raise ValueError("ไม่พบคอลัมน์ที่จำเป็น: " + ", ".join(missing))

    seen_ids: set[str] = set()
    valid_rows: list[dict] = []
    row_errors: list[dict] = []
    validator = _validate_feedback if kind == "feedback" else _validate_outcome
    for row_number, raw in rows:
        try:
            if kind == "feedback":
                value, errors = validator(raw, db, source_id, seen_ids)
            else:
                value, errors = validator(raw, db, seen_ids)
        except Exception:
            value, errors = None, ["เกิดข้อผิดพลาดระหว่างตรวจสอบแถวนี้"]
        if errors:
            row_errors.append({"row": row_number, "message": "；".join(errors)})
        elif value is not None:
            valid_rows.append(value)
    return valid_rows, row_errors, len(rows)


def process_import_batch(batch_id: int) -> None:
    """FastAPI BackgroundTasks entry point; opens its own DB session."""
    db = SessionLocal()
    storage_key = None
    try:
        batch = db.get(models.ImportBatch, batch_id)
        if batch is None or batch.status != "processing":
            return
        storage_key = batch.storage_key
        if not storage_key:
            raise ValueError("ไม่พบไฟล์ต้นฉบับสำหรับประมวลผล")
        data = storage.get_file(storage_key)
        valid, errors, total = validate_import(
            data, batch.file_name, batch.kind, db, batch.feedback_source_id
        )
        batch.total_rows = total
        batch.valid_rows = len(valid)
        batch.invalid_rows = len(errors)
        batch.errors = errors
        batch.preview_data = valid
        batch.status = "previewed"
        # Delete the raw upload before committing a preview that no longer needs it.
        storage.delete_file(storage_key)
        batch.storage_key = None
        db.commit()
        storage_key = None
    except Exception as exc:
        db.rollback()
        batch = db.get(models.ImportBatch, batch_id)
        if batch is not None:
            batch.status = "failed"
            batch.total_rows = 0
            batch.valid_rows = 0
            batch.invalid_rows = 0
            batch.preview_data = []
            batch.errors = [{"row": None, "message": str(exc) if isinstance(exc, ValueError)
                             else "ไม่สามารถประมวลผลไฟล์ได้ กรุณาตรวจสอบรูปแบบไฟล์"}]
            batch.storage_key = None
            db.commit()
    finally:
        if storage_key:
            try:
                storage.delete_file(storage_key)
            except Exception:
                # A failed cleanup is not allowed to hide the batch's validation result.
                pass
        db.close()
